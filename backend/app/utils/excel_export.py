from io import BytesIO
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from app.schemas.advertisement import AdvertisementResponse


_HEADER_FILL = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
_ALT_FILL = PatternFill(start_color="EBF0F7", end_color="EBF0F7", fill_type="solid")
_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_COLUMNS = [
    ("ID", 38),
    ("Title", 35),
    ("Media Type", 16),
    ("Status", 14),
    ("Plan", 14),
    ("Start Date", 22),
    ("End Date", 22),
    ("Created At", 22),
]


def generate_excel(advertisements: list[AdvertisementResponse]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Advertisements"

    # ── Title row ──────────────────────────────────────────────────
    ws.merge_cells(f"A1:{get_column_letter(len(_COLUMNS))}1")
    title_cell = ws["A1"]
    title_cell.value = f"Advertisement Report — {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
    title_cell.font = Font(bold=True, size=13, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="0D2137", end_color="0D2137", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # ── Header row ─────────────────────────────────────────────────
    for col_idx, (header, width) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[2].height = 22

    # ── Data rows ──────────────────────────────────────────────────
    for row_idx, ad in enumerate(advertisements, start=3):
        row_fill = _ALT_FILL if row_idx % 2 == 0 else None
        row_data = [
            str(ad.id),
            ad.title,
            ad.media_type.value,
            ad.status.value,
            ad.plan_type.value,
            ad.start_date.strftime("%Y-%m-%d %H:%M") if ad.start_date else "",
            ad.end_date.strftime("%Y-%m-%d %H:%M") if ad.end_date else "",
            ad.created_at.strftime("%Y-%m-%d %H:%M") if ad.created_at else "",
        ]
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if row_fill:
                cell.fill = row_fill
        ws.row_dimensions[row_idx].height = 18

    # ── Freeze header rows ────────────────────────────────────────
    ws.freeze_panes = "A3"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
