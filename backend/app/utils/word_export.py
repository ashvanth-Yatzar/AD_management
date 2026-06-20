from io import BytesIO
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches, Cm

from app.schemas.advertisement import AdvertisementResponse


def _set_cell_bg(cell, hex_color: str) -> None:
    """Helper to set table cell background colour."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def generate_word(advertisements: list[AdvertisementResponse]) -> BytesIO:
    doc = Document()

    # ── Page margins ───────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    # ── Document title ─────────────────────────────────────────────
    title_para = doc.add_heading("Advertisement Management Report", level=1)
    title_para.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph(
        f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}  |  "
        f"Total Records: {len(advertisements)}"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(9)
    subtitle.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    doc.add_paragraph()

    # ── Table ──────────────────────────────────────────────────────
    headers = ["#", "Title", "Media Type", "Status", "Plan", "Start Date", "End Date", "Created At"]
    col_widths = [Cm(1), Cm(5.5), Cm(2.8), Cm(2.2), Cm(2.8), Cm(3), Cm(3), Cm(3)]

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"

    # Header row
    hdr_row = table.rows[0]
    for idx, (header, width) in enumerate(zip(headers, col_widths)):
        cell = hdr_row.cells[idx]
        cell.width = width
        cell.text = header
        run = cell.paragraphs[0].runs[0]
        run.font.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_cell_bg(cell, "1E3A5F")

    # Data rows
    for row_idx, ad in enumerate(advertisements, start=1):
        row_data = [
            str(row_idx),
            ad.title,
            ad.media_type.value,
            ad.status.value.capitalize(),
            ad.plan_type.value.replace("_", " ").title(),
            ad.start_date.strftime("%Y-%m-%d") if ad.start_date else "",
            ad.end_date.strftime("%Y-%m-%d") if ad.end_date else "",
            ad.created_at.strftime("%Y-%m-%d") if ad.created_at else "",
        ]
        row = table.add_row()
        bg_color = "EBF0F7" if row_idx % 2 == 0 else "FFFFFF"

        for col_idx, value in enumerate(row_data):
            cell = row.cells[col_idx]
            cell.width = col_widths[col_idx]
            cell.text = value
            cell.paragraphs[0].runs[0].font.size = Pt(8)
            _set_cell_bg(cell, bg_color)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
